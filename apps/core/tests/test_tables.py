"""Downloads as CSV or XLSX: numbers stay numbers, nothing runs as a formula, lists keep their filters."""

import io
from datetime import date
from decimal import Decimal

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from apps.contacts.models import Contact
from apps.core.tables import ExportFormat, download
from apps.uploads.importers import IMPORTERS, Context, apply
from apps.uploads.services import PendingFile

ROWS = [["Name", "Amount", "Date"], ["=HYPERLINK(\"http://x\")", Decimal("1500.50"), date(2026, 1, 31)], ["PT B", Decimal("2000"), None]]


def test_csv_is_excel_friendly_and_neutralises_formulas():
    response = download("list", ROWS, ExportFormat.CSV)
    assert response["Content-Disposition"] == 'attachment; filename="list.csv"'
    text = response.content.decode("utf-8")
    assert text.startswith("﻿Name;Amount;Date") and "'=HYPERLINK" in text and "1500.50;2026-01-31" in text


def test_xlsx_keeps_numbers_and_dates_and_never_writes_a_formula():
    response = download("list", ROWS, ExportFormat.XLSX, title="Daftar")
    sheet = load_workbook(io.BytesIO(response.content)).active
    assert sheet.title == "Daftar" and sheet.freeze_panes == "A2"
    name, amount, on = (cell for cell in sheet[2])
    assert name.data_type == "s" and name.value.startswith("=HYPERLINK")  # text, not a formula
    assert amount.value == 1500.5 and amount.number_format == "#,##0.00" and on.value.date() == date(2026, 1, 31)


@pytest.mark.parametrize("export_format", ExportFormat.values)
def test_list_pages_download_their_filtered_rows(auth_client, organization, customer, supplier, export_format):
    response = auth_client.get(reverse("contacts:list"), {"kind": Contact.Kind.SUPPLIER, "download": export_format, "page": "9"})
    assert response.status_code == 200 and response["Content-Disposition"].endswith(f'contacts.{export_format}"')
    if export_format == ExportFormat.CSV:
        text = response.content.decode("utf-8-sig")
        assert "UD Pemasok" in text and "PT Pelanggan" not in text
    for name in ("invoicing:payments", "invoicing:items"):
        assert auth_client.get(reverse(name), {"download": export_format}).status_code == 200
    assert auth_client.get(reverse("invoicing:list", args=["sales"]), {"download": export_format}).status_code == 200
    assert "text/html" in auth_client.get(reverse("contacts:list"), {"download": "pdf"})["Content-Type"]  # unknown: the page


def test_downloaded_contacts_upload_again_unchanged(auth_client, organization, user, jatake):
    Contact.objects.create(organization=organization, name="PT Bolak Balik", npwp="012345678901000", address="Jl. Raya 1", village=jatake,
                           payment_terms_days=45)
    exported = auth_client.get(reverse("contacts:list"), {"download": "csv"}).content
    other = Contact.objects.get(name="PT Bolak Balik")
    other.delete()
    report = apply(IMPORTERS["contacts"], Context(organization, user, {}), PendingFile("contacts", "contacts.csv", exported, {}))
    assert report.ok, [(r.number, r.message) for r in report.results]
    again = Contact.objects.get(organization=organization, name="PT Bolak Balik")
    assert (again.npwp, again.village, again.payment_terms_days) == ("012345678901000", jatake, 45)


def test_xlsx_formats_follow_each_column_and_show_local_time():
    from datetime import datetime, timezone as dt_timezone

    rows = [["Amount", "Quantity", "Paid at"], [Decimal("2000.00"), Decimal("1.500"), datetime(2026, 10, 15, 3, 30, tzinfo=dt_timezone.utc)],
            [Decimal("1500.50"), Decimal("2.000"), None]]
    sheet = load_workbook(io.BytesIO(download("x", rows, ExportFormat.XLSX).content)).active
    (whole, quantity, paid), (cents, _q, _p) = (list(sheet[2]), list(sheet[3]))
    assert whole.number_format == cents.number_format == "#,##0.00"  # one format per money column
    assert quantity.number_format == "#,##0.000"
    assert paid.value == datetime(2026, 10, 15, 10, 30) and paid.number_format == "yyyy-mm-dd hh:mm"  # WIB, not UTC


def test_list_downloads_are_bounded(auth_client, customer, supplier, settings):
    url = reverse("contacts:list")
    settings.MAX_EXPORT_ROWS = 1
    response = auth_client.get(url, {"download": "csv", "q": ""})
    assert response.status_code == 302 and response.url == f"{url}?q="  # back to the list, with its filter
    settings.MAX_EXPORT_ROWS, settings.LIST_EXPORTS_PER_HOUR = 100, 1
    assert auth_client.get(url, {"download": "csv"}).status_code == 200
    assert auth_client.get(url, {"download": "csv"}).status_code == 302  # the hour's limit


def test_the_payments_download_reads_each_contact_with_its_payment(auth_client, organization, customer, account):
    from django.db import connection
    from django.test.utils import CaptureQueriesContext

    from apps.invoicing import services as invoicing
    from apps.invoicing.models import InvoiceKind
    from apps.invoicing.services import InvoiceLineInput

    for day in range(1, 6):
        invoice = invoicing.issue_invoice(invoice=invoicing.create_invoice(
            organization=organization, kind=InvoiceKind.SALES, contact=customer, date=date(2026, 3, day),
            lines=[InvoiceLineInput("Jasa", account("4-1100"), Decimal("1"), Decimal("1000"))],
        ))
        invoicing.record_payment(invoice=invoice, date=date(2026, 3, day), amount=Decimal("1000"), cash_account=account("1-1100"))
    with CaptureQueriesContext(connection) as queries:
        response = auth_client.get(reverse("invoicing:payments"), {"download": "csv"})
    assert response.status_code == 200 and response.content.decode("utf-8-sig").count(customer.name) == 5
    contact_lookups = [q["sql"] for q in queries.captured_queries if q["sql"].lstrip().startswith('SELECT "contacts_contact"')]
    assert not contact_lookups  # read with the payments (payments_for joins them), not once per payment
