"""Tables to download, as CSV or Excel (XLSX), from a header row and data rows.

CSV is what Excel opens in either language (UTF-8 with a BOM, ``;``-separated); XLSX keeps
amounts as numbers and dates as dates, with the header frozen. Text a spreadsheet would run as a
formula is neutralised in both (OWASP CSV injection): ``csv_safe`` in CSV, a text cell in XLSX.
"""

import csv
import io
from collections.abc import Iterable, Iterator
from datetime import date, datetime
from decimal import Decimal
from itertools import chain, islice

from django.db import models
from django.http import HttpResponse
from django.utils import timezone
from django.utils.translation import gettext_lazy as _

from .sanitizers import csv_safe

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_COLUMN_WIDTH = 60
WIDTH_SAMPLE_ROWS = 200  # column widths come from the first rows; the rest stream to the file


class ExportFormat(models.TextChoices):
    CSV = "csv", _("CSV")
    XLSX = "xlsx", _("Excel")


def _csv(rows: Iterable[list]) -> bytes:
    buffer = io.StringIO()
    buffer.write("\ufeff")  # BOM so Excel opens UTF-8 correctly
    writer = csv.writer(buffer, delimiter=";")
    writer.writerows([csv_safe(cell) for cell in row] for row in rows)
    return buffer.getvalue().encode("utf-8")


def _xlsx_value(value):
    """A cell's value: numbers and dates as such (datetimes in local time), anything else as text."""
    if isinstance(value, datetime):
        return timezone.localtime(value).replace(tzinfo=None) if timezone.is_aware(value) else value
    if isinstance(value, (int, float, Decimal, date)) and not isinstance(value, bool):
        return value
    return "" if value is None else str(value)


def _number_format(value) -> str | None:
    if isinstance(value, datetime):
        return "yyyy-mm-dd hh:mm"
    if isinstance(value, date):
        return "yyyy-mm-dd"
    if isinstance(value, Decimal):
        # The value's own scale is its column's (money 2 places, quantities 3), so a column is uniform.
        places = max(-value.as_tuple().exponent, 0) if value.is_finite() else 0
        return "#,##0" + ("." + "0" * places if places else "")
    return None


def _xlsx(rows: Iterable[list], title: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font

    book = Workbook(write_only=True)  # rows stream to the file instead of living in memory as cells
    sheet = book.create_sheet((title or "Sheet")[:31])
    rows = iter(rows)
    sample = [[_xlsx_value(value) for value in row] for row in islice(rows, WIDTH_SAMPLE_ROWS)]
    widths: dict[int, int] = {}
    for row in sample:
        for column, value in enumerate(row):
            widths[column] = max(widths.get(column, 0), min(MAX_COLUMN_WIDTH, len(str(value)) + 2))
    for column, width in widths.items():
        sheet.column_dimensions[_column_letter(column)].width = width
    sheet.freeze_panes = "A2"
    bold = Font(bold=True)
    remaining = ([_xlsx_value(value) for value in row] for row in rows)
    for row_number, row in enumerate(chain(sample, remaining), start=1):
        cells = []
        for value in row:
            cell = WriteOnlyCell(sheet, value=value)
            if isinstance(value, str):
                cell.data_type = "s"  # never a formula, whatever the text starts with
            elif number_format := _number_format(value):
                cell.number_format = number_format
            if row_number == 1:
                cell.font = bold
            cells.append(cell)
        sheet.append(cells)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _column_letter(index: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(index + 1)


def download(name: str, rows: Iterable[list] | Iterator[list], export_format: str, *, title: str = "") -> HttpResponse:
    """``name`` without its extension (``trial-balance``); ``rows`` start with the header."""
    export_format = ExportFormat(export_format)
    if export_format == ExportFormat.XLSX:
        response = HttpResponse(_xlsx(rows, title or name), content_type=XLSX_TYPE)
    else:
        response = HttpResponse(_csv(rows), content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{name}.{export_format.value}"'
    return response
