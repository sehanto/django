"""Reading uploaded spreadsheets (pure: no database access).

Uploads are CSV (UTF-8 or Windows-1252, comma, semicolon or tab separated, as Excel saves them in
either language) or XLSX (the first sheet, as Coretax and the banks download them). A table is
found by its header row, which may come after a few title lines; columns are matched by name in
English or Indonesian, in any order and case, so a file downloaded elsewhere needs no editing.
Values are read strictly: an amount, date or choice that can't be read is the row's error, never
a guess.
"""

import csv
import io
import re
import unicodedata
from collections.abc import Iterable, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from enum import StrEnum

from django.db import models
from django.utils.translation import gettext as _
from django.utils.translation import gettext_lazy

from apps.core.exceptions import ValidationFailed
from apps.core.sanitizers import sanitize_text

MAX_UPLOAD_BYTES = 2 * 1024 * 1024
MAX_ROWS = 5000
MAX_CELL_LENGTH = 2000
HEADER_SEARCH_ROWS = 15  # title lines a download may have above its header


class UploadKind(models.TextChoices):
    CONTACTS = "contacts", gettext_lazy("Contacts")
    ITEMS = "items", gettext_lazy("Items")
    ACCOUNTS = "accounts", gettext_lazy("Accounts")
    OPENING_BALANCES = "opening_balances", gettext_lazy("Opening balances")
    FIXED_ASSETS = "fixed_assets", gettext_lazy("Fixed assets brought over")
    OUTPUT_FAKTURS = "output_fakturs", gettext_lazy("Faktur Pajak numbers (Coretax, sales)")
    INPUT_FAKTURS = "input_fakturs", gettext_lazy("Input VAT fakturs (Coretax, purchases)")
    WITHHOLDING_SLIPS = "withholding_slips", gettext_lazy("Bukti potong numbers (Coretax)")
    BANK_STATEMENT = "bank_statement", gettext_lazy("Bank statement")


class RowOutcome(models.TextChoices):
    """What an upload does with a row."""

    CREATED = "created", gettext_lazy("Added")
    UPDATED = "updated", gettext_lazy("Recorded")
    SKIPPED = "skipped", gettext_lazy("Skipped")
    ERROR = "error", gettext_lazy("Error")


class FileFormat(StrEnum):
    CSV = "csv"
    XLSX = "xlsx"


@dataclass(frozen=True)
class Column:
    key: str
    label: str  # the header in this project's template
    aliases: tuple[str, ...] = ()  # other headers that mean the same, e.g. Coretax's or a bank's
    required: bool = False

    def names(self) -> set[str]:
        return {normalized_header(name) for name in (self.key, self.label, *self.aliases)}


@dataclass
class Row:
    number: int  # the line in the file, as a spreadsheet shows it
    values: dict[str, str]

    def get(self, key: str) -> str:
        return self.values.get(key, "")


@dataclass
class Table:
    rows: list[Row]
    missing: list[str] = field(default_factory=list)  # required columns not found
    found: set[str] = field(default_factory=set)  # the columns the file has


@dataclass
class RowResult:
    number: int
    outcome: str
    label: str = ""
    message: str = ""

    @property
    def outcome_label(self) -> str:
        return RowOutcome(self.outcome).label


@dataclass
class Report:
    results: list[RowResult] = field(default_factory=list)
    problems: list[str] = field(default_factory=list)  # about the file as a whole
    notes: list[str] = field(default_factory=list)  # e.g. documents the file didn't mention

    def add(self, number: int, outcome: str, label: str = "", message: str = "") -> None:
        self.results.append(RowResult(number, outcome, label, message))

    def count(self, outcome: str) -> int:
        return sum(1 for r in self.results if r.outcome == outcome)

    @property
    def errors(self) -> list[RowResult]:
        return [r for r in self.results if r.outcome == RowOutcome.ERROR]

    @property
    def ok(self) -> bool:
        return not self.problems and not self.errors

    @property
    def applied(self) -> int:
        return self.count(RowOutcome.CREATED) + self.count(RowOutcome.UPDATED)


def normalized_header(name: str) -> str:
    text = unicodedata.normalize("NFKD", str(name)).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "", text)


def file_format(name: str) -> FileFormat:
    suffix = name.rsplit(".", 1)[-1].lower() if "." in name else ""
    if suffix in ("csv", "txt"):
        return FileFormat.CSV
    if suffix == "xlsx":
        return FileFormat.XLSX
    raise ValidationFailed(_("Upload a CSV or XLSX file."))


def _decode(content: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValidationFailed(_("The file's characters can't be read: save it as CSV UTF-8."))


DELIMITERS = (",", ";", "\t", "|")  # Excel saves CSV with the list separator of its language (";" in Indonesian)


def _csv_rows(content: bytes, delimiter: str) -> list[list[str]]:
    text = _decode(content)
    if "\x00" in text:
        raise ValidationFailed(_("This isn't a CSV file."))
    try:
        return list(csv.reader(io.StringIO(text), delimiter=delimiter))
    except csv.Error as exc:
        raise ValidationFailed(_("The CSV file can't be read: %(error)s") % {"error": exc}) from None


def _header(raw: list[list[str]], columns: Sequence[Column]) -> tuple[int | None, dict[str, int]]:
    """The header row (the first that names every required column, else the one naming the most)."""
    by_name = {name: column.key for column in columns for name in column.names()}
    header_index, positions = None, {}
    for index, cells in enumerate(raw[:HEADER_SEARCH_ROWS]):
        found = {}
        for position, cell in enumerate(cells):
            key = by_name.get(normalized_header(cell))
            if key and key not in found:
                found[key] = position
        if len(found) > len(positions):
            header_index, positions = index, found
        if positions and all(c.key in positions for c in columns if c.required):
            break
    return header_index, positions


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if not (value.hour or value.minute or value.second) else value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        return f"{Decimal(str(value)).normalize():f}"
    return str(value)


def _xlsx_rows(content: bytes) -> list[list[str]]:
    import zipfile

    from openpyxl import (
        load_workbook,  # reads through defusedxml when installed (it is)
    )

    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:  # refuse zip bombs before parsing
            if sum(info.file_size for info in archive.infolist()) > 40 * MAX_UPLOAD_BYTES:
                raise ValidationFailed(_("The XLSX file is too large once unpacked."))
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except ValidationFailed:
        raise
    except Exception:  # noqa: BLE001 - openpyxl raises many kinds for a damaged or disguised file
        raise ValidationFailed(_("The XLSX file can't be read. Save it again from the spreadsheet app.")) from None
    try:
        sheet = workbook.worksheets[0]
        rows = []
        for row in sheet.iter_rows(values_only=True):
            rows.append([_cell(value) for value in row])
            if len(rows) > MAX_ROWS + HEADER_SEARCH_ROWS:
                break
        return rows
    finally:
        workbook.close()


def _clean_cell(value: str) -> str:
    text = sanitize_text(value)
    if text[:2] in ("'=", "'+", "'-", "'@"):  # this project's own exports neutralize formulas this way
        text = text[1:]
    return text


def read_table(content: bytes, name: str, columns: Sequence[Column]) -> Table:
    """The rows of an uploaded file under the given columns (unknown columns are ignored)."""
    if len(content) > MAX_UPLOAD_BYTES:
        raise ValidationFailed(_("The file is larger than 2 MB: split it into smaller files."))
    if not content.strip():
        raise ValidationFailed(_("The file is empty."))
    if file_format(name) == FileFormat.XLSX:
        raw = _xlsx_rows(content)
        header_index, positions = _header(raw, columns)
    else:
        best = max((_csv_rows(content, delimiter) for delimiter in DELIMITERS), key=lambda rows: len(_header(rows, columns)[1]))
        raw = best
        header_index, positions = _header(raw, columns)
    if header_index is None:
        raise ValidationFailed(_("No header row found: the first row names the columns, e.g. %(columns)s.") % {
            "columns": ", ".join(c.label for c in columns if c.required) or columns[0].label,
        })
    missing = [str(c.label) for c in columns if c.required and c.key not in positions]
    rows = []
    for offset, cells in enumerate(raw[header_index + 1:], start=header_index + 2):
        values = {key: _clean_cell(cells[position]) if position < len(cells) else "" for key, position in positions.items()}
        if not any(values.values()):
            continue
        if any(len(value) > MAX_CELL_LENGTH for value in values.values()):
            raise ValidationFailed(_("Row %(row)s has a cell longer than %(limit)s characters.") % {"row": offset, "limit": MAX_CELL_LENGTH})
        rows.append(Row(offset, values))
        if len(rows) > MAX_ROWS:
            raise ValidationFailed(_("The file has more than %(limit)s rows: split it into smaller files.") % {"limit": MAX_ROWS})
    if not rows and not missing:
        raise ValidationFailed(_("The file has no rows under its header."))
    return Table(rows, missing, set(positions))


# --- Values -----------------------------------------------------------------------

_AMOUNT = re.compile(r"^[+-]?\d+(?:[.,]\d+)*$")


def parse_amount(text: str, *, allow_negative: bool = False) -> Decimal | None:
    """An amount as people write it: ``1.234.567,89`` (Indonesian), ``1,234,567.89``, ``Rp 1.500``,
    ``(1.500)`` or ``-1500`` (negative). Empty is None."""
    value = text.strip().upper().replace("RP", "").replace("IDR", "").replace(" ", "").replace(" ", "")
    if not value:
        return None
    negative = value.startswith("(") and value.endswith(")")
    if negative:
        value = value[1:-1]
    if value.startswith("-"):
        negative, value = True, value[1:]
    elif value.startswith("+"):
        value = value[1:]
    if not _AMOUNT.match(value):
        raise ValueError(text)
    last_dot, last_comma = value.rfind("."), value.rfind(",")
    decimal_mark = max(last_dot, last_comma)
    if decimal_mark != -1:
        tail = value[decimal_mark + 1:]
        separator = value[decimal_mark]
        # One separator followed by exactly three digits is a thousands separator (1.500 / 1,500).
        thousands_only = len(tail) == 3 and value.count(separator) >= 1 and (
            (separator == "." and last_comma == -1) or (separator == "," and last_dot == -1)
        )
        if thousands_only:
            value = value.replace(separator, "")
        else:
            whole = re.sub(r"[.,]", "", value[:decimal_mark])
            value = f"{whole}.{tail}"
    try:
        amount = Decimal(value)
    except InvalidOperation:
        raise ValueError(text) from None
    if negative:
        if not allow_negative:
            raise ValueError(text)
        amount = -amount
    return amount


def parse_rate(text: str) -> Decimal | None:
    """A percentage such as ``2``, ``1,5`` or ``0.5`` (``%`` allowed). Empty is None."""
    value = text.strip().rstrip("%").strip().replace(",", ".")
    if not value:
        return None
    try:
        return Decimal(value)
    except InvalidOperation:
        raise ValueError(text) from None


MONTHS = {
    "jan": 1, "januari": 1, "january": 1, "feb": 2, "februari": 2, "february": 2, "peb": 2, "mar": 3, "maret": 3, "march": 3,
    "apr": 4, "april": 4, "mei": 5, "may": 5, "jun": 6, "juni": 6, "june": 6, "jul": 7, "juli": 7, "july": 7,
    "agu": 8, "agt": 8, "ags": 8, "agustus": 8, "aug": 8, "august": 8, "sep": 9, "sept": 9, "september": 9,
    "okt": 10, "oktober": 10, "oct": 10, "october": 10, "nov": 11, "nop": 11, "november": 11, "nopember": 11,
    "des": 12, "desember": 12, "dec": 12, "december": 12,
}


def parse_date(text: str) -> date | None:
    """``2026-01-31`` (ISO, as XLSX cells read), ``31/01/2026``, ``31-01-2026``, ``31.01.2026`` or
    ``31 Jan 2026`` / ``31 Januari 2026`` (day first, as written in Indonesia). Empty is None."""
    value = text.strip()
    if not value:
        return None
    value = value.split(" ")[0] if re.match(r"^\d{4}-\d{2}-\d{2}[ T]\d", value) else value
    match = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})$", value)
    if match:
        year, month, day = map(int, match.groups())
    elif match := re.match(r"^(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})$", value):
        day, month, year = map(int, match.groups())
    elif match := re.match(r"^(\d{1,2})[\s-]+([A-Za-z]+)[\s-]+(\d{4})$", value):
        day, name, year = int(match[1]), match[2].lower(), int(match[3])
        if name not in MONTHS:
            raise ValueError(text)
        month = MONTHS[name]
    else:
        raise ValueError(text)
    return date(year, month, day)


YES = {"1", "y", "ya", "yes", "true", "benar", "v", "x", "aktif"}
NO = {"", "0", "n", "no", "tidak", "t", "false", "salah", "nonaktif"}


def parse_bool(text: str, *, default: bool = False) -> bool:
    value = text.strip().lower()
    if not value:
        return default
    if value in YES:
        return True
    if value in NO:
        return False
    raise ValueError(text)


def parse_choice(text: str, choices: Iterable[tuple[str, Iterable[str]]]) -> str | None:
    """The value whose names (its code and labels, in any language) match ``text``; None if none."""
    wanted = normalized_header(text)
    for value, names in choices:
        if wanted in {normalized_header(value), *(normalized_header(name) for name in names)}:
            return value
    return None


def digits(text: str) -> str:
    return re.sub(r"\D", "", text)
